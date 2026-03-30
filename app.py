from flask import Flask, render_template, request, redirect, send_file
import json
from datetime import datetime
from docx import Document

app = Flask(__name__)

# -------------------------------
# MASTER DATA
# -------------------------------

namelist = [
    "Nirmallya Shil","Narayan Ch. Shil","Gopal Ch. Sarkar"
]

addresslist = [
    "Jyotinagar, Phansidewa",
    "Jyotinagar, Phansidewa",
    "Jyotinagar, Phansidewa"
]

familycode = [
    "4417144700","5193400","1159917000"
]

# Make family codes 12 digits
familycode = [code.zfill(12) for code in familycode]

# -------------------------------
# LOAD & SAVE DATA
# -------------------------------

def load_data():
    try:
        with open("data.json", "r") as f:
            return json.load(f)
    except:
        return []

def save_data(data):
    with open("data.json", "w") as f:
        json.dump(data, f, indent=4)

# -------------------------------
# HOME
# -------------------------------

@app.route("/")
def home():
    return render_template("home.html")

# -------------------------------
# CREATE LIST
# -------------------------------

@app.route("/create", methods=["GET","POST"])
def create():

    if request.method == "POST":

        amounts = []
        total = 0

        for i in range(len(namelist)):

            amt = request.form.get(namelist[i])

            if not amt or not amt.isdigit():
                amt = 0
            else:
                amt = int(amt)

            amounts.append(amt)
            total += amt

        data = load_data()

        new_entry = {
            "date": datetime.now().strftime("%d-%m-%Y %H:%M"),
            "names": namelist,
            "addresses": addresslist,
            "familycode": familycode,
            "amounts": amounts,
            "total": total
        }

        data.append(new_entry)
        save_data(data)

        return redirect("/storage")

    return render_template("create.html", names=namelist)

# -------------------------------
# STORAGE
# -------------------------------

@app.route("/storage")
def storage():
    data = load_data()
    return render_template("storage.html", data=data)

# -------------------------------
# DOWNLOAD DOCX
# -------------------------------

@app.route("/download/<int:index>")
def download(index):

    data = load_data()
    item = data[index]

    doc = Document()

    # HEADER
    doc.add_heading("SATSANG PHILANTHROPY, SATSANG DEOGHAR", 0)

    doc.add_paragraph("JYOTINAGAR, PHANSIDEWA UPOYOJANA KENDRA-43")
    doc.add_paragraph("TOTAL AMOUNT - __________________________")
    doc.add_paragraph("(Rupees __________________________________ only)")
    doc.add_paragraph("POWER JYOTI CHALLAN, SBI, RANIDANGA")
    doc.add_paragraph("JOURNAL NO. - __________________     Dtd - __________________")
    doc.add_paragraph("Start Srl No - 01     End Srl No - ____")
    doc.add_paragraph("Deposited By: NITYANANDA SHIL")
    doc.add_paragraph("Family Code: 000005193400")

    doc.add_paragraph(" ")

    # TABLE
    table = doc.add_table(rows=1, cols=5)

    header = table.rows[0].cells
    header[0].text = "Srl No"
    header[1].text = "Name"
    header[2].text = "Address"
    header[3].text = "Family Code"
    header[4].text = "Amount"

    for i in range(len(item["names"])):

        row = table.add_row().cells

        row[0].text = str(i+1)
        row[1].text = item["names"][i]

        # SAME AS ABOVE LOGIC
        if i > 0 and item["addresses"][i] == item["addresses"][i-1]:
            row[2].text = '"'
        else:
            row[2].text = item["addresses"][i]

        row[3].text = item["familycode"][i]
        row[4].text = str(item["amounts"][i])

    # FOOTER
    doc.add_paragraph("\nTOTAL AMOUNT = " + str(item["total"]))

    # SAVE FILE
    file_path = "satsang_list.docx"
    doc.save(file_path)

    return send_file(file_path, as_attachment=True)

# -------------------------------
# DELETE LIST
# -------------------------------

@app.route("/delete/<int:index>")
def delete(index):

    data = load_data()

    if 0 <= index < len(data):
        data.pop(index)

    save_data(data)

    return redirect("/storage")

# -------------------------------
# RUN APP
# -------------------------------

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)